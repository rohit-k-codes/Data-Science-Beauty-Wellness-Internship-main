"""Week 5 models for the dataset's observed recommendation indicator.

Run from the repository root:
    python src/model.py

By default, the pipeline uses a reproducible class-proportional sample from
the full Week 3 cleaned file. Preprocessing is fitted inside each training
fold. Rating and post-review feedback fields are deliberately excluded from
predictors because they can directly reveal or post-date the target.
"""

import argparse
import json
import logging
from pathlib import Path
from typing import Any, Optional

import matplotlib

matplotlib.use("Agg")

import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
import seaborn as sns
import sklearn
from sklearn.compose import ColumnTransformer
from sklearn.dummy import DummyClassifier
from sklearn.feature_extraction.text import TfidfVectorizer
from sklearn.linear_model import LogisticRegression
from sklearn.metrics import (
    accuracy_score,
    confusion_matrix,
    f1_score,
    precision_score,
    recall_score,
    roc_auc_score,
    roc_curve,
)
from sklearn.model_selection import GroupKFold, GroupShuffleSplit, cross_validate
from sklearn.pipeline import Pipeline
from sklearn.preprocessing import OneHotEncoder
from sklearn.tree import DecisionTreeClassifier


ROOT = Path(__file__).resolve().parent.parent
DEFAULT_INPUT = ROOT / "data" / "processed" / "reviews_clean.csv"
DEFAULT_SUMMARY = ROOT / "data" / "processed" / "model_summary.json"
DEFAULT_CONFUSION = ROOT / "figures" / "model_confusion_matrices.png"
DEFAULT_ROC = ROOT / "figures" / "model_roc_curves.png"
DEFAULT_REPORT = ROOT / "reports" / "Week_5_Report.docx"

TARGET = "is_recommended"
TEXT_FEATURE = "review_text_clean"
PROFILE_FEATURE = "skin_type"
GROUP_COLUMN = "author_id"
REQUIRED_COLUMNS = {
    TARGET,
    "rating",
    TEXT_FEATURE,
    PROFILE_FEATURE,
    GROUP_COLUMN,
}
PREDICTOR_COLUMNS = [TEXT_FEATURE, PROFILE_FEATURE]
RANDOM_STATE = 42
DEFAULT_SAMPLE_SIZE = 30_000
DEFAULT_CHUNK_SIZE = 50_000
DEFAULT_CV_FOLDS = 3
MAX_TEXT_FEATURES = 12_000

log = logging.getLogger(__name__)


def _input_schema(input_path: Path) -> list[str]:
    if not input_path.is_file():
        raise FileNotFoundError(f"Processed dataset not found: {input_path}")
    header = pd.read_csv(input_path, nrows=0)
    missing = REQUIRED_COLUMNS - set(header.columns)
    if missing:
        raise ValueError(
            "Required modeling columns are missing: "
            f"{sorted(missing)}. Available columns: {list(header.columns)}"
        )
    return header.columns.tolist()


def collect_labeled_sample(
    input_path: Path,
    sample_size: int = DEFAULT_SAMPLE_SIZE,
    chunk_size: int = DEFAULT_CHUNK_SIZE,
    random_state: int = RANDOM_STATE,
) -> tuple[pd.DataFrame, dict[str, Any]]:
    """Stream the cleaned dataset and select a class-proportional uniform sample.

    Independent random priorities implement reservoir-style uniform sampling
    within each valid target class. The full file is read to calculate exact
    label/missingness counts, but only a bounded sample is retained in memory.
    """
    if sample_size < 6:
        raise ValueError("sample_size must be at least 6 for a three-fold split.")
    if chunk_size <= 0:
        raise ValueError("chunk_size must be a positive integer.")
    input_columns = _input_schema(input_path)
    rng = np.random.default_rng(random_state)

    rows_seen = 0
    missing_counts: dict[str, int] = {}
    target_counts: dict[str, int] = {}
    rating_target_crosstab: dict[str, dict[str, int]] = {}
    input_dtypes: dict[str, str] = {}
    reservoirs: dict[int, pd.DataFrame] = {
        0: pd.DataFrame(),
        1: pd.DataFrame(),
    }

    usecols = [TARGET, TEXT_FEATURE, PROFILE_FEATURE, GROUP_COLUMN, "rating"]
    for chunk in pd.read_csv(
        input_path,
        usecols=usecols,
        chunksize=chunk_size,
        low_memory=False,
    ):
        rows_seen += len(chunk)
        if not input_dtypes:
            input_dtypes = {
                str(column): str(dtype) for column, dtype in chunk.dtypes.items()
            }
        for column, count in chunk.isna().sum().items():
            missing_counts[column] = missing_counts.get(column, 0) + int(count)

        numeric_target = pd.to_numeric(chunk[TARGET], errors="coerce")
        target_valid = numeric_target.isin([0, 1])
        for label, count in numeric_target[target_valid].value_counts().items():
            label_int = int(label)
            target_counts[str(label_int)] = (
                target_counts.get(str(label_int), 0) + int(count)
            )
        crosstab = pd.crosstab(chunk["rating"], numeric_target)
        for rating_value, row in crosstab.iterrows():
            rating_key = str(rating_value)
            rating_target_crosstab.setdefault(rating_key, {})
            for target_value, count in row.items():
                target_key = str(int(target_value))
                rating_target_crosstab[rating_key][target_key] = (
                    rating_target_crosstab[rating_key].get(target_key, 0)
                    + int(count)
                )

        eligible = chunk.loc[target_valid].copy()
        eligible[TARGET] = numeric_target.loc[target_valid].astype("int8")
        eligible[TEXT_FEATURE] = eligible[TEXT_FEATURE].fillna("").astype(str)
        eligible[PROFILE_FEATURE] = (
            eligible[PROFILE_FEATURE].fillna("unknown").astype(str).str.strip()
        )
        eligible[GROUP_COLUMN] = eligible[GROUP_COLUMN].astype("string")
        eligible = eligible.loc[
            eligible[TEXT_FEATURE].str.strip().ne("")
            & eligible[GROUP_COLUMN].notna()
            & eligible[GROUP_COLUMN].str.strip().ne("")
        ]

        for label in (0, 1):
            class_rows = eligible.loc[eligible[TARGET].eq(label)].copy()
            if class_rows.empty:
                continue
            class_rows["_sample_priority"] = rng.random(len(class_rows))
            candidates = pd.concat(
                [reservoirs[label], class_rows],
                ignore_index=True,
            )
            reservoirs[label] = candidates.nsmallest(
                min(sample_size, len(candidates)),
                columns="_sample_priority",
            ).reset_index(drop=True)

    valid_class_counts = {
        label: target_counts.get(str(label), 0) for label in (0, 1)
    }
    labeled_total = sum(valid_class_counts.values())
    if labeled_total == 0:
        raise ValueError("No valid 0/1 recommendation targets found in input.")

    requested_total = min(sample_size, labeled_total)
    negative_share = valid_class_counts[0] / labeled_total
    negative_n = int(round(requested_total * negative_share))
    positive_n = requested_total - negative_n
    negative_n = min(negative_n, len(reservoirs[0]))
    positive_n = min(positive_n, len(reservoirs[1]))
    selected = []
    if negative_n:
        selected.append(reservoirs[0].nsmallest(negative_n, "_sample_priority"))
    if positive_n:
        selected.append(reservoirs[1].nsmallest(positive_n, "_sample_priority"))
    if not selected:
        raise ValueError(
            "No labeled rows with non-empty text and reviewer IDs are available."
        )

    sample = pd.concat(selected, ignore_index=True)
    sample = sample.drop(columns="_sample_priority")
    sample[TARGET] = sample[TARGET].astype("int8")
    sample[PROFILE_FEATURE] = sample[PROFILE_FEATURE].replace("", "unknown")
    sample = sample.sample(frac=1, random_state=random_state).reset_index(drop=True)
    sample_counts = {
        str(label): int(count)
        for label, count in sample[TARGET].value_counts().sort_index().items()
    }
    metadata = {
        "input_rows": int(rows_seen),
        "input_columns": input_columns,
        "input_dtypes": input_dtypes,
        "input_missing_values": missing_counts,
        "target_counts": target_counts,
        "rating_vs_recommendation_counts": rating_target_crosstab,
        "invalid_or_missing_target_rows": int(
            rows_seen - sum(valid_class_counts.values())
        ),
        "valid_labeled_rows": int(labeled_total),
        "sample_rows": int(len(sample)),
        "sample_target_counts": sample_counts,
        "sample_unique_reviewers": int(sample[GROUP_COLUMN].nunique()),
        "sample_method": (
            "Uniform priority-reservoir sample within each target class, "
            "allocated proportionally to full-file class counts."
        ),
        "random_state": random_state,
    }
    return sample, metadata


def prepare_xy(
    sample: pd.DataFrame,
) -> tuple[pd.DataFrame, pd.Series, pd.Series]:
    """Select safe predictors, binary outcome, and reviewer grouping labels."""
    missing = REQUIRED_COLUMNS - set(sample.columns)
    if missing:
        raise ValueError(f"Modeling sample is missing required columns: {sorted(missing)}")

    target = pd.to_numeric(sample[TARGET], errors="coerce")
    valid = target.isin([0, 1])
    text = sample[TEXT_FEATURE].fillna("").astype(str)
    groups = sample[GROUP_COLUMN].astype("string")
    valid &= text.str.strip().ne("") & groups.notna() & groups.str.strip().ne("")

    x = pd.DataFrame(
        {
            TEXT_FEATURE: text.loc[valid],
            PROFILE_FEATURE: (
                sample.loc[valid, PROFILE_FEATURE]
                .fillna("unknown")
                .astype(str)
                .str.strip()
                .replace("", "unknown")
            ),
        },
        index=sample.index[valid],
    )
    y = target.loc[valid].astype("int8").rename(TARGET)
    return x, y, groups.loc[valid].astype(str)


def build_models(
    max_text_features: int = MAX_TEXT_FEATURES,
) -> dict[str, Pipeline]:
    """Build comparable text/profile pipelines with train-only preprocessing."""
    if max_text_features <= 0:
        raise ValueError("max_text_features must be positive.")

    def feature_transformer() -> ColumnTransformer:
        return ColumnTransformer(
            transformers=[
                (
                    "text",
                    TfidfVectorizer(
                        max_features=max_text_features,
                        min_df=2,
                        ngram_range=(1, 2),
                        sublinear_tf=True,
                        strip_accents="unicode",
                    ),
                    TEXT_FEATURE,
                ),
                (
                    "skin_type",
                    OneHotEncoder(handle_unknown="ignore"),
                    [PROFILE_FEATURE],
                ),
            ],
            remainder="drop",
        )

    return {
        "logistic_regression": Pipeline(
            steps=[
                ("features", feature_transformer()),
                (
                    "classifier",
                    LogisticRegression(
                        class_weight="balanced",
                        max_iter=300,
                        random_state=RANDOM_STATE,
                        solver="liblinear",
                    ),
                ),
            ]
        ),
        "decision_tree": Pipeline(
            steps=[
                ("features", feature_transformer()),
                (
                    "classifier",
                    DecisionTreeClassifier(
                        class_weight="balanced",
                        max_depth=20,
                        min_samples_leaf=20,
                        random_state=RANDOM_STATE,
                    ),
                ),
            ]
        ),
    }


def _positive_probability(estimator: Any, x: pd.DataFrame) -> np.ndarray:
    probabilities = estimator.predict_proba(x)
    classes = list(estimator.classes_)
    if 1 not in classes:
        return np.zeros(len(x), dtype=float)
    return probabilities[:, classes.index(1)]


def calculate_metrics(
    y_true: pd.Series | np.ndarray,
    y_pred: np.ndarray,
    y_probability: Optional[np.ndarray] = None,
) -> dict[str, Optional[float]]:
    """Calculate held-out binary classification metrics for recommended=1."""
    metrics: dict[str, Optional[float]] = {
        "accuracy": float(accuracy_score(y_true, y_pred)),
        "precision_positive": float(
            precision_score(y_true, y_pred, pos_label=1, zero_division=0)
        ),
        "recall_positive": float(
            recall_score(y_true, y_pred, pos_label=1, zero_division=0)
        ),
        "f1_positive": float(f1_score(y_true, y_pred, pos_label=1, zero_division=0)),
        "f1_macro": float(f1_score(y_true, y_pred, average="macro", zero_division=0)),
        "roc_auc": None,
    }
    if y_probability is not None and len(np.unique(y_true)) == 2:
        metrics["roc_auc"] = float(roc_auc_score(y_true, y_probability))
    return metrics


def _cross_validation(
    models: dict[str, Pipeline],
    x_train: pd.DataFrame,
    y_train: pd.Series,
    groups_train: pd.Series,
    folds: int,
) -> dict[str, dict[str, dict[str, Optional[float]]]]:
    """Evaluate model pipelines using reviewer-grouped cross-validation."""
    group_count = int(groups_train.nunique())
    split_count = min(folds, group_count)
    if split_count < 2:
        raise ValueError(
            "At least two distinct reviewer groups are needed for cross-validation."
        )
    group_cv = GroupKFold(n_splits=split_count)
    scores = {
        "accuracy": "accuracy",
        "precision_positive": "precision",
        "recall_positive": "recall",
        "f1_positive": "f1",
        "f1_macro": "f1_macro",
        "roc_auc": "roc_auc",
    }
    output: dict[str, dict[str, dict[str, Optional[float]]]] = {}
    for name, model in models.items():
        result = cross_validate(
            model,
            x_train,
            y_train,
            groups=groups_train,
            cv=group_cv,
            scoring=scores,
            n_jobs=1,
            error_score="raise",
            return_train_score=True,
        )
        model_scores: dict[str, dict[str, Optional[float]]] = {}
        for metric in scores:
            test_scores = result[f"test_{metric}"]
            train_scores = result[f"train_{metric}"]
            model_scores[metric] = {
                "test_mean": float(np.mean(test_scores)),
                "test_std": float(np.std(test_scores, ddof=1)),
                "train_mean": float(np.mean(train_scores)),
            }
        model_scores["fit_time_seconds"] = {
            "test_mean": float(np.mean(result["fit_time"])),
            "test_std": float(np.std(result["fit_time"], ddof=1)),
            "train_mean": None,
        }
        output[name] = model_scores
    return output


def _safe_save_figure(fig: plt.Figure, path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    fig.tight_layout()
    fig.savefig(path, dpi=160, bbox_inches="tight")
    plt.close(fig)
    log.info("Figure saved: %s", path)


def plot_evaluation(
    y_test: pd.Series,
    model_outputs: dict[str, dict[str, Any]],
    confusion_path: Path,
    roc_path: Path,
) -> None:
    """Plot held-out confusion matrices and ROC curves from actual predictions."""
    sns.set_theme(style="whitegrid")
    names = list(model_outputs)
    fig, axes = plt.subplots(1, len(names), figsize=(5 * len(names), 4.5))
    if len(names) == 1:
        axes = [axes]
    for axis, name in zip(axes, names):
        matrix = confusion_matrix(y_test, model_outputs[name]["predictions"], labels=[0, 1])
        sns.heatmap(
            matrix,
            annot=True,
            fmt="d",
            cmap="Blues",
            cbar=False,
            xticklabels=["Not recommended (0)", "Recommended (1)"],
            yticklabels=["Not recommended (0)", "Recommended (1)"],
            ax=axis,
        )
        axis.set_title(name.replace("_", " ").title())
        axis.set_xlabel("Predicted")
        axis.set_ylabel("Actual")
        axis.tick_params(axis="x", labelrotation=20)
    _safe_save_figure(fig, confusion_path)

    fig, ax = plt.subplots(figsize=(7, 5.5))
    for name in names:
        probabilities = model_outputs[name]["probabilities"]
        fpr, tpr, _ = roc_curve(y_test, probabilities)
        auc = model_outputs[name]["metrics"]["roc_auc"]
        ax.plot(fpr, tpr, label=f"{name.replace('_', ' ').title()} (AUC={auc:.3f})")
    ax.plot([0, 1], [0, 1], linestyle="--", color="gray", label="Chance (AUC=0.500)")
    ax.set(
        title=f"ROC curves on reviewer-held-out test set (n={len(y_test):,})",
        xlabel="False positive rate",
        ylabel="True positive rate",
        xlim=(0, 1),
        ylim=(0, 1),
    )
    ax.legend(loc="lower right")
    _safe_save_figure(fig, roc_path)


def write_report(summary: dict[str, Any], report_path: Path) -> None:
    """Write a Word report with measured modeling results and caveats."""
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor

    document = Document()
    document.styles["Normal"].font.name = "Arial"
    document.styles["Normal"].font.size = Pt(10)
    for style_name in ("Heading 1", "Heading 2"):
        document.styles[style_name].font.name = "Arial"
        document.styles[style_name].font.color.rgb = RGBColor(31, 78, 121)

    title = document.add_heading(
        "Week 5: Predictive Modeling and Algorithm Selection", 0
    )
    title.alignment = 1
    subtitle = document.add_paragraph(
        "Analyzing Customer Reviews for Personalized Skincare Recommendations"
    )
    subtitle.alignment = 1
    subtitle.runs[0].bold = True
    document.add_paragraph("Beauty & Wellness Data Science Internship").alignment = 1

    target_counts = summary["data"]["target_counts"]
    total = summary["data"]["input_rows"]
    labeled = summary["data"]["valid_labeled_rows"]
    test_size = summary["split"]["test_rows"]
    model_metrics = summary["test_metrics"]
    baseline = model_metrics["most_frequent_baseline"]
    best_name = max(
        [name for name in model_metrics if name != "most_frequent_baseline"],
        key=lambda name: model_metrics[name]["f1_macro"],
    )

    document.add_heading("1. Modeling objective and research-question connection", level=1)
    document.add_paragraph(
        "Predict the dataset's observed is_recommended flag (0/1) from the review's "
        "cleaned text and the reviewer's self-reported skin_type. This is a "
        "secondary predictive task related to whether review content and profile "
        "information are associated with a recorded recommendation label. It "
        "does not extract or validate skin concerns and does not predict whether "
        "a product is suitable for an individual."
    )
    document.add_paragraph(
        "The task does not directly answer RQ2 (whether different skin profiles "
        "rate the same product differently). It is an exploratory supervised "
        "proxy using the dataset's recommendation outcome."
    )

    document.add_heading("2. Dataset and target variable", level=1)
    document.add_paragraph(
        f"Input: {summary['data']['input_file']}. The full cleaned file contains "
        f"{total:,} rows. The binary target is available for {labeled:,} rows: "
        f"{target_counts.get('0', 0):,} not recommended and "
        f"{target_counts.get('1', 0):,} recommended. "
        f"{summary['data']['invalid_or_missing_target_rows']:,} rows have missing "
        "or invalid target values and were excluded from training/evaluation."
    )
    document.add_paragraph(
        f"The modeling sample contains {summary['data']['sample_rows']:,} labeled "
        f"rows ({summary['data']['sample_target_counts']}). Sampling: "
        f"{summary['data']['sample_method']}"
    )
    document.add_paragraph(
        f"Input missingness among modeling columns: "
        f"{summary['data']['input_missing_values']}."
    )

    document.add_heading("3. Feature selection and preprocessing", level=1)
    document.add_paragraph(
        "Predictors: review_text_clean transformed with word-level TF-IDF "
        "(unigrams and bigrams; max_features=12,000; min_df=2) and skin_type "
        "transformed with one-hot encoding. Empty/missing profile values are "
        "represented as unknown. The preprocessing steps are inside each "
        "scikit-learn Pipeline and are fitted only on each training fold."
    )
    document.add_paragraph(
        "Leakage controls: rating was excluded because it is strongly associated "
        "with is_recommended in the source (rating is part of the outcome "
        "generation process). is_recommended itself, rating_band, helpfulness, "
        "feedback counts, product IDs/names, author IDs, and other post-outcome "
        "or identifier-like columns were not predictors. author_id was used only "
        "to keep reviewers in a single split/fold."
    )

    document.add_heading("4. Models and rationale", level=1)
    document.add_paragraph(
        "Baseline: DummyClassifier(strategy='most_frequent'), which predicts the "
        "training-set majority class. Logistic Regression with balanced class "
        "weights provides a regularized linear model over sparse TF-IDF/profile "
        "features. A depth-limited Decision Tree with a minimum leaf size tests "
        "a nonlinear tree-based alternative while controlling complexity."
    )

    document.add_heading("5. Baseline and validation strategy", level=1)
    document.add_paragraph(
        f"A reviewer-grouped holdout used GroupShuffleSplit with "
        f"{summary['split']['train_rows']:,} training and {test_size:,} test "
        f"observations. No author appears in both partitions. The test set "
        f"contains {summary['split']['test_unique_reviewers']:,} reviewers; the "
        f"training set contains {summary['split']['train_unique_reviewers']:,}. "
        f"Three-fold GroupKFold cross-validation was performed on training data, "
        f"grouped by reviewer. Sampling and splitting used random_state="
        f"{summary['data']['random_state']}."
    )
    document.add_paragraph(
        "Reported metrics: accuracy, positive-class precision/recall/F1, macro "
        "F1, and ROC-AUC. Positive class means is_recommended=1. The majority "
        "baseline's ROC-AUC is 0.500 because it assigns the same score to all "
        "observations."
    )

    document.add_heading("6. Actual held-out evaluation", level=1)
    table = document.add_table(rows=1, cols=7)
    table.style = "Light Shading Accent 1"
    headers = ["Model", "Accuracy", "Precision", "Recall", "F1 positive", "Macro F1", "ROC-AUC"]
    for cell, value in zip(table.rows[0].cells, headers):
        cell.text = value
    for name, values in model_metrics.items():
        cells = table.add_row().cells
        cells[0].text = name.replace("_", " ").title()
        for cell, key in zip(cells[1:], [
            "accuracy", "precision_positive", "recall_positive",
            "f1_positive", "f1_macro", "roc_auc",
        ]):
            cell.text = f"{values[key]:.3f}" if values[key] is not None else "N/A"

    document.add_paragraph(
        f"On this held-out sample, {best_name.replace('_', ' ').title()} had the "
        f"highest macro F1 ({model_metrics[best_name]['f1_macro']:.3f}). The "
        f"majority baseline accuracy was {baseline['accuracy']:.3f}; because "
        "class 1 is common, accuracy alone is not sufficient."
    )

    document.add_heading("7. Cross-validation and fit diagnostics", level=1)
    cv_table = document.add_table(rows=1, cols=4)
    cv_table.style = "Light Shading Accent 1"
    for cell, value in zip(cv_table.rows[0].cells, ["Model", "CV macro F1 mean ± SD", "Train macro F1", "CV ROC-AUC mean ± SD"]):
        cell.text = value
    for name, values in summary["cross_validation"].items():
        cells = cv_table.add_row().cells
        cells[0].text = name.replace("_", " ").title()
        cells[1].text = f"{values['f1_macro']['test_mean']:.3f} ± {values['f1_macro']['test_std']:.3f}"
        cells[2].text = f"{values['f1_macro']['train_mean']:.3f}"
        cells[3].text = f"{values['roc_auc']['test_mean']:.3f} ± {values['roc_auc']['test_std']:.3f}"
    document.add_paragraph(
        "Training-versus-cross-validation macro F1 provides a limited "
        "overfitting diagnostic. A large gap suggests overfitting; low training "
        "and validation scores suggest underfitting. These diagnostics apply to "
        "the sampled data and selected features only."
    )

    document.add_heading("8. Evaluation plots", level=1)
    for key, caption in [
        ("confusion_matrix_plot", "Held-out confusion matrices; rows are actual labels and columns are predicted labels."),
        ("roc_plot", "Held-out receiver operating characteristic curves, computed from predicted positive-class probabilities."),
    ]:
        document.add_picture(summary["outputs"][key], width=Inches(6.1))
        document.add_paragraph(caption, style="Caption")

    document.add_heading("9. Business interpretation", level=1)
    document.add_paragraph(
        f"The best macro-F1 model ({best_name.replace('_', ' ')}) can be "
        "interpreted only as predicting the platform's review-level "
        "recommendation flag from review wording and a self-reported profile. "
        "Text may encode a reviewer's expressed opinion. This is not evidence "
        "that skin type causes satisfaction or that the product is clinically "
        "or personally suitable. The baseline comparison and minority-class "
        "recall/precision should guide any further evaluation, not a direct "
        "consumer recommendation."
    )

    document.add_heading("10. Limitations and Week 6 recommendations", level=1)
    limitations = [
        "The evaluation uses a 30,000-row class-proportional sample of the full cleaned dataset, not all 811,974 labeled examples.",
        "is_recommended is strongly associated with rating; rating was excluded to prevent target leakage. Review text may still express or reveal sentiment closely related to the label.",
        "Reviewer-grouped splitting prevents the same reviewer from appearing in train and test; product IDs are excluded from predictors, but products may occur in both partitions.",
        "The target is an observed review-level flag, not an expert assessment of efficacy, skin concerns, safety, or suitability.",
        "The class distribution is imbalanced; accuracy can overstate useful performance. Compare macro F1, positive precision/recall, ROC-AUC, and the baseline.",
        "Cross-validation and holdout results depend on the fixed sample/split and may not generalize across time, products, or other platforms.",
        "Week 6 should audit errors by skin-type availability and product, calibrate scores if needed, test temporal/product-group splits, quantify uncertainty, and evaluate fairness and business utility before any recommendation use.",
    ]
    for item in limitations:
        document.add_paragraph(item, style="List Bullet")

    document.add_heading("Conclusion", level=1)
    document.add_paragraph(
        f"Actual results show whether review text plus self-reported skin type "
        f"predict the dataset's recommendation indicator better than the "
        f"majority baseline on a reviewer-held-out sample. The highest macro-F1 "
        f"model was {best_name.replace('_', ' ')} at "
        f"{model_metrics[best_name]['f1_macro']:.3f}. This is an exploratory "
        "benchmark only, not a validated personalized skincare recommender. "
        "Expected outcomes were not substituted for measured metrics."
    )

    document.add_heading("References", level=1)
    for reference in [
        "Sephora Products and Skincare Reviews dataset, Kaggle: https://www.kaggle.com/datasets/nadyinky/sephora-products-and-skincare-reviews (local processed file inspected; license not verified).",
        "scikit-learn documentation: https://scikit-learn.org/stable/",
        "scikit-learn Pipeline and ColumnTransformer documentation: https://scikit-learn.org/stable/modules/compose.html",
        "scikit-learn model evaluation documentation: https://scikit-learn.org/stable/modules/model_evaluation.html",
        "pandas documentation: https://pandas.pydata.org/docs/",
    ]:
        document.add_paragraph(reference, style="List Bullet")

    report_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(report_path)
    log.info("Week 5 report saved: %s", report_path)


def run_modeling(
    input_path: Path = DEFAULT_INPUT,
    summary_path: Path = DEFAULT_SUMMARY,
    confusion_path: Path = DEFAULT_CONFUSION,
    roc_path: Path = DEFAULT_ROC,
    report_path: Path = DEFAULT_REPORT,
    sample_size: int = DEFAULT_SAMPLE_SIZE,
    chunk_size: int = DEFAULT_CHUNK_SIZE,
    cv_folds: int = DEFAULT_CV_FOLDS,
    max_text_features: int = MAX_TEXT_FEATURES,
    overwrite: bool = False,
) -> dict[str, Any]:
    """Train and evaluate baseline and classifiers on the cleaned review file."""
    if cv_folds < 2:
        raise ValueError("cv_folds must be at least 2.")
    output_paths = [summary_path, confusion_path, roc_path, report_path]
    existing = [path for path in output_paths if path.exists()]
    if existing and not overwrite:
        raise FileExistsError(
            "Refusing to overwrite existing modeling output(s); pass --overwrite: "
            + ", ".join(str(path) for path in existing)
        )

    sample, data_metadata = collect_labeled_sample(
        input_path=input_path,
        sample_size=sample_size,
        chunk_size=chunk_size,
    )
    x, y, groups = prepare_xy(sample)
    if y.nunique() < 2:
        raise ValueError("The selected modeling sample must contain both target classes.")

    splitter = GroupShuffleSplit(
        n_splits=1,
        test_size=0.2,
        random_state=RANDOM_STATE,
    )
    train_indices, test_indices = next(splitter.split(x, y, groups=groups))
    x_train, x_test = x.iloc[train_indices], x.iloc[test_indices]
    y_train, y_test = y.iloc[train_indices], y.iloc[test_indices]
    groups_train, groups_test = groups.iloc[train_indices], groups.iloc[test_indices]
    overlap = set(groups_train).intersection(groups_test)
    if overlap:
        raise RuntimeError(
            f"Reviewer leakage detected: {len(overlap)} groups cross the holdout."
        )
    if y_train.nunique() < 2 or y_test.nunique() < 2:
        raise ValueError(
            "Both target classes must appear in train and test partitions; "
            "increase --sample-size."
        )

    baseline = DummyClassifier(strategy="most_frequent")
    baseline.fit(np.zeros((len(y_train), 1)), y_train)
    baseline_pred = baseline.predict(np.zeros((len(y_test), 1)))
    baseline_probability = baseline.predict_proba(np.zeros((len(y_test), 1)))[:, 1]
    test_metrics: dict[str, dict[str, Any]] = {
        "most_frequent_baseline": calculate_metrics(
            y_test,
            baseline_pred,
            baseline_probability,
        )
    }

    model_outputs: dict[str, dict[str, Any]] = {
        "most_frequent_baseline": {
            "predictions": baseline_pred,
            "probabilities": baseline_probability,
            "metrics": test_metrics["most_frequent_baseline"],
        }
    }
    models = build_models(max_text_features=max_text_features)
    for name, model in models.items():
        log.info("Fitting %s on %d reviewer-grouped training rows", name, len(y_train))
        model.fit(x_train, y_train)
        predictions = model.predict(x_test)
        probabilities = _positive_probability(model, x_test)
        metrics = calculate_metrics(y_test, predictions, probabilities)
        test_metrics[name] = metrics
        model_outputs[name] = {
            "predictions": predictions,
            "probabilities": probabilities,
            "metrics": metrics,
        }

    cross_validation = _cross_validation(
        models,
        x_train,
        y_train,
        groups_train,
        folds=cv_folds,
    )
    plot_evaluation(y_test, model_outputs, confusion_path, roc_path)

    summary = {
        "task": {
            "name": "Review-level recommendation indicator classification",
            "target": TARGET,
            "positive_class": 1,
            "research_question_connection": (
                "Exploratory proxy for whether review text and reported profile "
                "are associated with the dataset's review-level recommendation "
                "flag; not a product suitability or skin-concern model."
            ),
        },
        "data": {
            **data_metadata,
            "input_file": str(input_path),
            "predictor_columns": PREDICTOR_COLUMNS,
            "excluded_columns": [
                "rating",
                TARGET,
                "rating_band",
                "helpfulness",
                "total_feedback_count",
                "total_neg_feedback_count",
                "total_pos_feedback_count",
                "product_id",
                "product_name",
                GROUP_COLUMN,
            ],
            "modeling_sample_target_counts": {
                str(label): int(count)
                for label, count in y.value_counts().sort_index().items()
            },
        },
        "split": {
            "method": "GroupShuffleSplit by author_id; test_size=0.20",
            "random_state": RANDOM_STATE,
            "train_rows": int(len(y_train)),
            "test_rows": int(len(y_test)),
            "train_class_counts": {
                str(label): int(count)
                for label, count in y_train.value_counts().sort_index().items()
            },
            "test_class_counts": {
                str(label): int(count)
                for label, count in y_test.value_counts().sort_index().items()
            },
            "train_unique_reviewers": int(groups_train.nunique()),
            "test_unique_reviewers": int(groups_test.nunique()),
            "reviewer_overlap": int(len(overlap)),
        },
        "models": {
            "most_frequent_baseline": "DummyClassifier(strategy='most_frequent')",
            "logistic_regression": (
                "TF-IDF word unigrams/bigrams + skin_type one-hot; "
                "LogisticRegression(class_weight='balanced', solver='liblinear')"
            ),
            "decision_tree": (
                "TF-IDF word unigrams/bigrams + skin_type one-hot; "
                "DecisionTreeClassifier(class_weight='balanced', max_depth=20, "
                "min_samples_leaf=20)"
            ),
            "max_text_features": max_text_features,
            "vectorizer_min_df": 2,
            "class_weight": "balanced for logistic regression and decision tree",
        },
        "test_metrics": test_metrics,
        "cross_validation": cross_validation,
        "software": {
            "scikit_learn": sklearn.__version__,
            "pandas": pd.__version__,
            "numpy": np.__version__,
            "random_state": RANDOM_STATE,
            "group_cv_folds": min(cv_folds, int(groups_train.nunique())),
        },
        "outputs": {
            "confusion_matrix_plot": str(confusion_path),
            "roc_plot": str(roc_path),
            "summary_json": str(summary_path),
            "report_docx": str(report_path),
        },
        "interpretation_limits": [
            "The target is the dataset's review-level is_recommended label.",
            "No rating, rating-derived feature, product identifier, or reviewer identifier was used as a predictor.",
            "The evaluation is sample-based and is not evidence of product efficacy or individual suitability.",
            "Review text may itself express sentiment related to the target.",
        ],
    }

    summary_path.parent.mkdir(parents=True, exist_ok=True)
    with summary_path.open("w", encoding="utf-8") as output:
        json.dump(summary, output, indent=2, ensure_ascii=False, allow_nan=False)
    write_report(summary, report_path)

    print(f"Input rows: {data_metadata['input_rows']:,}")
    print(f"Valid labeled rows: {data_metadata['valid_labeled_rows']:,}")
    print(f"Target counts (0/1): {data_metadata['target_counts']}")
    print(f"Sample rows: {len(y):,}; train: {len(y_train):,}; test: {len(y_test):,}")
    print(f"Reviewer overlap across holdout: {len(overlap)}")
    print("Held-out metrics:")
    print(json.dumps(test_metrics, indent=2))
    print("Reviewer-grouped cross-validation:")
    print(json.dumps(cross_validation, indent=2))
    print(f"Summary: {summary_path}")
    print(f"Report: {report_path}")
    return summary


def _parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Week 5 recommendation-label modeling pipeline."
    )
    parser.add_argument("--input", type=Path, default=DEFAULT_INPUT)
    parser.add_argument("--summary", type=Path, default=DEFAULT_SUMMARY)
    parser.add_argument("--confusion-plot", type=Path, default=DEFAULT_CONFUSION)
    parser.add_argument("--roc-plot", type=Path, default=DEFAULT_ROC)
    parser.add_argument("--report", type=Path, default=DEFAULT_REPORT)
    parser.add_argument("--sample-size", type=int, default=DEFAULT_SAMPLE_SIZE)
    parser.add_argument("--chunk-size", type=int, default=DEFAULT_CHUNK_SIZE)
    parser.add_argument("--cv-folds", type=int, default=DEFAULT_CV_FOLDS)
    parser.add_argument("--max-text-features", type=int, default=MAX_TEXT_FEATURES)
    parser.add_argument("--overwrite", action="store_true")
    return parser.parse_args()


if __name__ == "__main__":
    logging.basicConfig(level=logging.INFO, format="%(levelname)s %(message)s")
    args = _parse_args()
    run_modeling(
        input_path=args.input,
        summary_path=args.summary,
        confusion_path=args.confusion_plot,
        roc_path=args.roc_plot,
        report_path=args.report,
        sample_size=args.sample_size,
        chunk_size=args.chunk_size,
        cv_folds=args.cv_folds,
        max_text_features=args.max_text_features,
        overwrite=args.overwrite,
    )
